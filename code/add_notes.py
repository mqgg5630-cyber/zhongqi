#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""add_notes.py - 把 docs/ppt_outline.json 里的演讲备注写进任意 pptx（按页序对应）。

用法：
    python code/add_notes.py deliverable/中期答辩.pptx                 # 就地写入
    python code/add_notes.py in.pptx -o out.pptx --outline docs/ppt_outline.json

用途：A 版 PPT 由 make_ppt2.py 生成时没有备注；此脚本可把大纲里的
「演讲备注（口播稿）」补齐，三版口径一致。也可用于给别的 deck 批量补备注。
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", default=None, help="默认就地覆盖")
    ap.add_argument("--outline", default="docs/ppt_outline.json")
    args = ap.parse_args()

    data = json.loads(Path(args.outline).read_text(encoding="utf-8"))
    notes = [(s.get("note") or "").strip() for s in data["slides"]]

    prs = Presentation(args.pptx)
    n = min(len(notes), len(prs.slides._sldIdLst))
    filled = 0
    for slide, note in zip(prs.slides, notes[:n]):
        if not note:
            continue
        if slide.notes_slide.notes_text_frame.text.strip() == note:
            continue
        slide.notes_slide.notes_text_frame.text = note
        filled += 1

    out = Path(args.out or args.pptx)
    prs.save(str(out))
    print(f"{args.pptx}: 写入 {filled} 页演讲备注（共 {len(prs.slides._sldIdLst)} 页），"
          f"大纲提供 {sum(1 for x in notes if x)} 条 -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
