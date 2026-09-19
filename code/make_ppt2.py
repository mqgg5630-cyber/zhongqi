#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_ppt2.py - 中期答辩 PPT（A 版：学术蓝，讲思路与完成度）

本版的内容口径
----------------
* 只讲"研究思路 + 工作完成到哪一步"，不展开技术细节（页面上不出现参数、版本、命令、指标数字）
* 三模型（Attention / LSTM / BERT）共识预测按"已完成"呈现
* 分析完抗菌肽差异后，用宏蛋白组做二次去重，筛出 AD 组特有的（AD 特异性）抗菌肽
* 最后与 AD 发病机制建立关联（Aβ 聚集 / AChE–PAS / 免疫与炎症），并补一个抑菌实验验证
* 例外：工作进度一页保留百分比，用于说明"完成到哪一步"

输出：deliverable/中期答辩_A_学术蓝.pptx（16 页，16:9，全篇最小字号 15 pt）
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import make_ppt as mp          # 复用版式与自检工具

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FIG = mp.FIG
OUT = mp.OUT  # -> deliverable/中期答辩_A_学术蓝.pptx
MARGIN, CONTENT_W = mp.MARGIN, mp.CONTENT_W
CONTENT_TOP, CONTENT_BOTTOM, TAKEAWAY_TOP = mp.CONTENT_TOP, mp.CONTENT_BOTTOM, mp.TAKEAWAY_TOP
INK, BLUE, TEAL, ORANGE, GREY, LIGHT, SAND, MINT = (
    mp.INK, mp.BLUE, mp.TEAL, mp.ORANGE, mp.GREY, mp.LIGHT, mp.SAND, mp.MINT)

add_title = mp.add_title
textbox = mp.textbox
add_para = mp.add_para
add_rich = mp.add_rich
add_round_box = mp.add_round_box
add_takeaway = mp.add_takeaway
add_footer = mp.add_footer
place_figure = mp.place_figure
number_cards = mp.number_cards


def bullet_block(slide, bullets, left, top, width, size=17, gap=15, height=None):
    tb, tf = textbox(slide, left, top, width, height or (CONTENT_BOTTOM - top))
    for k, text in enumerate(bullets):
        add_rich(tf, [("▪　", False, BLUE), (text, False, INK)], size=size,
                 first=(k == 0), space_before=(0 if k == 0 else gap), space_after=0,
                 line_spacing=1.25)
    return tb


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = mp.SLIDE_W, mp.SLIDE_H
    slides = []

    # 1 ─ 封面 -------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN + Inches(3.5), Inches(2.16),
                            CONTENT_W - Inches(7.0), Pt(2.6))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background(); ln.shadow.inherit = False
    tb, tf = textbox(s, MARGIN, Inches(1.06), CONTENT_W, Inches(0.5))
    add_para(tf, "研究生论文中期检查", size=19, bold=True, color=BLUE, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0)
    tb, tf = textbox(s, MARGIN, Inches(2.46), CONTENT_W, Inches(1.7))
    add_para(tf, "基于深度学习的阿尔茨海默症患者与健康人群", size=30, bold=True,
             color=INK, first=True, align=PP_ALIGN.CENTER, space_before=0, space_after=8,
             line_spacing=1.15)
    add_para(tf, "肠道微生物组中抗菌肽的差异性研究", size=30, bold=True, color=INK,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0, line_spacing=1.15)
    tb, tf = textbox(s, MARGIN, Inches(4.62), CONTENT_W, Inches(1.9))
    for line in ["汇报人：文绍华　　学号：2024110316",
                 "指导教师：申亮　　培养单位：生命科学学院",
                 "学科专业：生物学（生物化学与分子生物学）",
                 "汇报日期：2026 年 9 月"]:
        add_para(tf, line, size=17, color=INK, align=PP_ALIGN.CENTER, space_before=10,
                 space_after=0)
    slides.append(s)

    # 2 ─ 汇报提纲 ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "汇报提纲")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP - Inches(0.05), CONTENT_W, Inches(4.3))
    items = [
        ("01", "研究思路：从肠道微生物组到候选抗菌肽"),
        ("02", "已完成：数据资源、短肽库与三模型共识预测"),
        ("03", "已完成：分阶段差异分析与宏蛋白组二次去重"),
        ("04", "下一步：与 AD 发病机制的关联分析"),
        ("05", "下一步：候选抗菌肽抑菌实验验证"),
        ("06", "工作进度与后续安排"),
    ]
    for k, (num, text) in enumerate(items):
        add_rich(tf, [(num + "　", True, TEAL), (text, False, INK)], size=20,
                 first=(k == 0), space_before=(0 if k == 0 else 20), space_after=0)
    add_takeaway(s, "本次汇报重点说明研究思路与各项工作的完成程度")
    slides.append(s)

    # 3 ─ 研究思路总览 -----------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "研究思路：八步走的整体框架与完成进度", "研究思路")
    place_figure(s, "figA_研究思路总览.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "主线：数据资源 → 短肽库 → 三模型共识预测 → 分阶段差异 → 去重筛选 AD 特有肽 → AD 特有肽与 AD 的关联分析")
    slides.append(s)

    # 4 ─ 科学问题 ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "科学问题：AD 患者与健康人群的抗菌肽是否存在差异", "研究背景与科学问题")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP, CONTENT_W, Inches(3.2))
    rows = [
        ("背景", "AD 的发生发展与肠—脑轴介导的慢性神经炎症密切相关，AD 患者与健康人群的"
                 "肠道微生物组在组成与功能上均存在差异。"),
        ("切入点", "微生物基因组中的小开放阅读框可编码短肽，其中包含具有抗菌与免疫调节功能的"
                   "抗菌肽，可能是连接微生物刺激与神经炎症的潜在效应分子。"),
        ("空白", "已有抗菌肽研究多基于通用训练集与宿主来源序列，尚缺少面向 AD 人群、"
                 "并按认知功能阶段分层比较的微生物源抗菌肽研究。"),
    ]
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "　", True, BLUE), (text, False, INK)], size=17,
                 first=(k == 0), space_before=(0 if k == 0 else 16), space_after=0,
                 line_spacing=1.25)
    box = add_round_box(s, MARGIN, Inches(5.0), CONTENT_W, Inches(1.16), fill=SAND, line=ORANGE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    add_para(tf, "科学问题", size=17, bold=True, color=ORANGE, first=True,
             space_before=0, space_after=4)
    add_para(tf, "AD 患者与健康人群肠道微生物组所编码的抗菌肽是否存在差异？"
                 "能否筛出 AD 组特有的抗菌肽，并与 AD 的发病机制建立联系？",
             size=18, bold=True, color=INK, space_before=0, space_after=0, line_spacing=1.15)
    slides.append(s)

    # 5 ─ 数据资源 ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成：数据资源与微生物源短肽库构建", "已完成工作 · 第一步")
    bullets = [
        "完成全队列宏基因组数据的统一处理与质量评估，获得高质量的微生物基因组参考集",
        "基于参考集完成小开放阅读框预测，构建覆盖全队列、非冗余的微生物源短肽库",
        "短肽库保留完整溯源信息，每条短肽均可回溯到对应基因组与临床样本",
        "参考集按物种多样性原则构建，减少因组装质量差异带来的样本偏倚",
    ]
    bullet_block(s, bullets, MARGIN, CONTENT_TOP + Inches(0.15), Inches(7.05))
    number_cards(s, [
        ("队列", "例完整宏基因组样本", BLUE),
        ("基因组", "参考集（含高质量基因组）", BLUE),
        ("短肽库", "非冗余微生物源短肽", TEAL),
        ("可溯源", "每条短肽可回溯至样本", TEAL),
    ], left=MARGIN + Inches(7.35), top=CONTENT_TOP + Inches(0.15),
        width=Inches(4.74), height=Inches(4.35))
    add_takeaway(s, "第一步与第二步均已完成：数据资源与短肽库已就绪，为后续预测与差异分析提供统一输入")
    slides.append(s)

    # 6 ─ 三模型共识预测 ---------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成：三模型共识预测候选抗菌肽", "已完成工作 · 第三步")
    place_figure(s, "figB_三模型预测.png", top=CONTENT_TOP - Inches(0.05),
                 max_h=Inches(3.35))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.05), CONTENT_W, Inches(1.0))
    add_para(tf, "采用 Attention、LSTM、BERT 三个模型对短肽库进行独立预测，"
                 "三者一致判为阳性的序列才纳入候选抗菌肽集合；"
                 "该步骤已在本阶段完成，得到候选抗菌肽名单。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "三模型协同提高候选集合可信度；该步骤已完成，候选抗菌肽名单已产出")
    slides.append(s)

    # 7 ─ 分阶段差异分析 ---------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成：分阶段抗菌肽差异分析", "已完成工作 · 第四步")
    place_figure(s, "figC_分阶段差异分析.png", top=CONTENT_TOP - Inches(0.05),
                 max_h=Inches(3.35))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.05), CONTENT_W, Inches(1.0))
    add_para(tf, "按认知阶段对队列分组，比较各阶段与健康人群之间候选抗菌肽的丰度与组成差异，"
                 "筛选出随病程变化明显的候选抗菌肽。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "先得到随病程变化的候选抗菌肽，再用表达证据过滤，为 AD 特异性特有肽筛选做准备")
    slides.append(s)

    # 8 ─ 宏蛋白组去重 -----------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成：宏蛋白组二次去重，筛选 AD 特异性特有抗菌肽", "已完成工作 · 第五步")
    place_figure(s, "figD_宏蛋白组去重.png", top=CONTENT_TOP - Inches(0.05),
                 max_h=Inches(3.35))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.05), CONTENT_W, Inches(1.0))
    add_para(tf, "在第一层序列去冗余之后，引入宏蛋白组表达证据进行二次去重，"
                 "只保留真实存在且被检出的抗菌肽，并据此筛出 AD 组特有（AD 特异性）的抗菌肽。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "二次去重解决两类问题：序列冗余，以及“有预测、无表达”的假阳性")
    slides.append(s)

    # 9 ─ 机制关联 ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "下一步：AD 特有肽与 AD 的关联分析（致病方向）", "下一步计划 · 第六步")
    place_figure(s, "figE_机制关联.png", top=CONTENT_TOP - Inches(0.05),
                 max_h=Inches(3.4))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.02), CONTENT_W, Inches(1.0))
    add_para(tf, "参照乙酰胆碱酯酶—β-淀粉样肽复合物分子模拟研究的思路，按致病方向考察 AD 特异性特有肽："
                 "与 Aβ 的相互作用是否促进成核与聚集，与 AChE 外周阴离子位点的结合是否推动成核环节，"
                 "以及在免疫与炎症通路上是否放大信号。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "关联分析的落点：判断 AD 特有肽是否推动 Aβ 聚集与神经炎症，而不是去证明它抑制 AD")
    slides.append(s)

    # 10 ─ 抑菌实验 --------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "下一步：候选抗菌肽抑菌实验验证", "下一步计划 · 第七步")
    place_figure(s, "figF_抑菌实验方案.png", top=CONTENT_TOP - Inches(0.02),
                 max_h=Inches(3.3))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.0), CONTENT_W, Inches(1.0))
    add_para(tf, "选取有代表性的候选抗菌肽进行人工合成，以大肠杆菌与金黄色葡萄球菌作为指示菌，"
                 "先以纸片扩散法观察抑菌圈进行初筛，再以微量肉汤稀释法测定最低抑菌浓度，"
                 "并设置阳性与阴性对照，以获得抑菌活性的直接证据。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "以通用、可及的微生物学方法，完成候选抗菌肽抑菌活性的直接验证")
    slides.append(s)

    # 11 ─ 进度 ------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "工作进度：主要分析已完成，进入验证与撰写阶段", "工作进度")
    place_figure(s, "figG_进度甘特.png", top=CONTENT_TOP - Inches(0.05),
                 max_h=Inches(3.35))
    tb, tf = textbox(s, MARGIN, CONTENT_BOTTOM - Inches(1.05), CONTENT_W, Inches(1.0))
    add_para(tf, "数据资源、短肽库、三模型共识预测、分阶段差异分析与宏蛋白组去重均已完成；"
                 "当前处于机制关联与抑菌实验验证阶段，论文撰写同步推进。",
             size=16.5, color=INK, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "前五项已完成，机制关联分析与抑菌实验验证正在推进，整体进度与开题计划一致")
    slides.append(s)

    # 12 ─ 后续安排 --------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "后续安排：四项收尾工作", "下一步计划")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, Inches(4.2))
    rows = [
        ("1. 机制关联分析", "完成 AD 特有肽与 Aβ 聚集、AChE 结合及炎症通路的关联分析。"),
        ("2. 抑菌实验验证", "完成候选抗菌肽的合成与抑菌实验，获得抑菌活性初筛结果。"),
        ("3. 结果整理", "整理 AD 特有肽、关联分析与验证结果，形成完整的图表与结论。"),
        ("4. 论文撰写", "完成学位论文撰写与投稿论文准备，按计划申请预答辩。"),
    ]
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "　", True, ORANGE), (text, False, INK)], size=17,
                 first=(k == 0), space_before=(0 if k == 0 else 18), space_after=0,
                 line_spacing=1.25)
    add_takeaway(s, "收尾工作集中在分析结论与实验验证，不再涉及大规模数据生产")
    slides.append(s)

    # 13 ─ 与开题的对应 ----------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "与开题计划的对应关系", "工作对照")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP + Inches(0.05), CONTENT_W, Inches(4.3))
    rows = [
        ("数据准备", "开题计划使用公开宏基因组数据 → 已完成全队列数据处理与参考集构建"),
        ("候选肽挖掘", "开题计划提取 sORF 并预测抗菌肽 → 已完成短肽库构建与三模型共识预测"),
        ("差异分析", "开题计划按病程阶段比较 → 已完成分阶段差异分析，并新增宏蛋白组二次去重"),
        ("模型方案", "开题拟构建 DeepMetaAMP → 实际采用 Attention / LSTM / BERT 三个已发表模型协同预测"),
        ("机制与验证", "开题拟做功能与可视化分析 → 下一步做 AD 特有肽与 AD 的关联分析与抑菌实验验证"),
    ]
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "　", True, TEAL), (text, False, INK)], size=16.5,
                 first=(k == 0), space_before=(0 if k == 0 else 16), space_after=0,
                 line_spacing=1.25)
    add_takeaway(s, "研究方向未变，仅在预测模型与去重环节做了更务实的调整")
    slides.append(s)

    # 14 ─ 问题与对策 ------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "存在问题与应对措施", "问题与对策")
    cards = [
        ("问题 1", "预测结果存在假阳性风险",
         "应对：以三模型共识判定提高候选集合可信度，再用宏蛋白组表达证据二次去重，最后以抑菌实验验证活性。"),
        ("问题 2", "部分样本缺少可用的参考基因组",
         "应对：在丰度分析阶段改用全量基因组回贴，恢复样本覆盖，避免参考集筛选造成样本缺失。"),
        ("问题 3", "机制关联的证据强度有限",
         "应对：以对接与分子模拟、文献比对为主，按致病方向检验候选肽是否推动 Aβ 聚集与神经炎症，"
         "明确研究定位为“线索发现”，不夸大因果结论。"),
    ]
    y = CONTENT_TOP + Inches(0.05)
    for tag, prob, fix in cards:
        box = add_round_box(s, MARGIN, y, CONTENT_W, Inches(1.28), fill=SAND, line=ORANGE)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.2)
        add_para(tf, f"{tag}　{prob}", size=17, bold=True, color=INK, first=True,
                 space_before=0, space_after=4, line_spacing=1.12)
        add_para(tf, fix, size=16.5, color=INK, space_before=0, space_after=0, line_spacing=1.12)
        y += Inches(1.42)
    add_takeaway(s, "三个问题都有具体对策，不影响整体进度")
    slides.append(s)

    # 15 ─ 预期成果 --------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "预期成果", "成果形式")
    bullets = [
        "一套完整的肠道微生物源抗菌肽挖掘分析流程，可复用于其他队列与疾病",
        "AD 特异性特有的候选抗菌肽清单，作为后续研究的线索集",
        "AD 特有肽与 AD 发病机制关联的分析结果（Aβ 聚集、AChE 结合、炎症通路）",
        "候选抗菌肽抑菌活性的初步实验证据",
        "以第一作者撰写并投稿 SCI 收录论文 1 篇，完成学位论文",
    ]
    bullet_block(s, bullets, MARGIN, CONTENT_TOP + Inches(0.15), CONTENT_W, size=17.5, gap=18)
    add_takeaway(s, "成果既包括数据与清单，也包括机制线索与实验证据")
    slides.append(s)

    # 16 ─ 致谢 ------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb, tf = textbox(s, MARGIN, Inches(2.6), CONTENT_W, Inches(1.5))
    add_para(tf, "恳请各位老师批评指正", size=34, bold=True, color=INK, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=14)
    add_para(tf, "谢谢！", size=22, color=BLUE, align=PP_ALIGN.CENTER, space_before=0,
             space_after=0)
    tb, tf = textbox(s, MARGIN, Inches(4.6), CONTENT_W, Inches(0.8))
    add_para(tf, "汇报人：文绍华　指导教师：申亮", size=16, color=GREY, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0)
    slides.append(s)

    for i, sl in enumerate(slides, 1):
        if i not in (1, len(slides)):
            add_footer(sl, i, len(slides))
    return prs


def main() -> int:
    if not (FIG / "figA_研究思路总览.png").exists():
        print("figures missing - run python code/make_figures2.py first", file=sys.stderr)
        return 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(str(OUT))
    print(f"wrote {OUT}")
    for p in mp.PROBLEMS:
        print("WARN ", p)
    return mp.audit(OUT)


if __name__ == "__main__":
    raise SystemExit(main())
