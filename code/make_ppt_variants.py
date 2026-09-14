#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_ppt_variants.py - 由 docs/ppt_outline.json 生成两种**另版风格**的中期答辩 PPT。

三版并列供选择（都是原生可编辑 pptx，真实文本框/形状，非图片化）：

  A  学术蓝（现有）      deliverable/中期答辩.pptx
  B  极简线框（本脚本）  deliverable/versions/中期答辩_B_极简线框.pptx
  C  卡片色块（本脚本）  deliverable/versions/中期答辩_C_卡片色块.pptx

三版内容完全一致（同一份 JSON），差别只在版式与配色；每页都写入大纲里的
「演讲备注」，可直接用 ppt-master 之类工具生成配音。

用法：
    python code/make_ppt_variants.py            # 生成 B、C 两版
    python code/make_ppt_variants.py --style B
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import make_ppt as mp

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUTLINE = Path("docs/ppt_outline.json")
OUT_DIR = Path("deliverable/versions")

SLIDE_W, SLIDE_H = mp.SLIDE_W, mp.SLIDE_H
MARGIN = mp.MARGIN
CONTENT_W = mp.CONTENT_W
MIN_PT = 15

INK = mp.INK
BLUE = mp.BLUE
TEAL = mp.TEAL
ORANGE = mp.ORANGE
GREY = mp.GREY
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ----------------------------------------------------------------- 主题定义
STYLES = {
    "B": {
        "name": "极简线框",
        "accent": RGBColor(0x0F, 0x4C, 0x5C),      # 深青
        "accent2": RGBColor(0x2F, 0x9E, 0x8F),
        "rule": RGBColor(0xD5, 0xDE, 0xE5),
        "ink": RGBColor(0x1A, 0x1F, 0x24),
        "muted": RGBColor(0x70, 0x7A, 0x85),
        "card_fill": None,                          # 极简：卡片不填充
        "card_line": RGBColor(0xC3, 0xCF, 0xD8),
        "header_band": None,                        # 无标题色带
        "side_bar": None,
        "takeaway_style": "rule",                   # 左侧竖线 + 文字
        "toc_number_size": 22,
        "title_size": 27,
        "body_size": 17,
    },
    "C": {
        "name": "卡片色块",
        "accent": RGBColor(0x14, 0x3C, 0x6B),      # 深蓝
        "accent2": RGBColor(0x2F, 0x6F, 0xB0),
        "rule": RGBColor(0xDA, 0xE4, 0xEE),
        "ink": RGBColor(0x12, 0x32, 0x4F),
        "muted": RGBColor(0x5E, 0x6E, 0x7E),
        "card_fill": RGBColor(0xF2, 0xF6, 0xFA),
        "card_line": RGBColor(0xC9, 0xD8, 0xE6),
        "header_band": RGBColor(0xE8, 0xEF, 0xF7),
        "side_bar": RGBColor(0x14, 0x3C, 0x6B),
        "takeaway_style": "band",
        "toc_number_size": 24,
        "title_size": 27,
        "body_size": 17,
    },
}

TITLE_TOP = Inches(0.34)
SUB_TOP = Inches(1.12)
BODY_TOP = Inches(1.62)
TAKE_TOP = Inches(6.30)
TAKE_H = Inches(0.60)
FOOT_TOP = Inches(7.04)


# ----------------------------------------------------------------- 绘制工具
def tb(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    return mp.textbox(slide, left, top, width, height, anchor)


def para(tf, text, size=17, bold=False, color=INK, first=False, before=6, after=2,
         align=PP_ALIGN.LEFT, spacing=1.22):
    return mp.add_para(tf, text, size=size, bold=bold, color=color, first=first,
                       space_before=before, space_after=after, align=align,
                       line_spacing=spacing)


def rich(tf, parts, size=17, first=False, before=6, after=0, spacing=1.25):
    return mp.add_rich(tf, parts, size=size, space_before=before, space_after=after,
                       first=first, line_spacing=spacing)


def hline(slide, left, top, width, color, weight=Pt(1.0)):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, weight)
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False
    return ln


def card(slide, left, top, width, height, style, fill=None, line=None, radius=0.07):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    f = style["card_fill"] if fill is None else fill
    if f is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = f
    sh.line.color.rgb = style["card_line"] if line is None else line
    sh.line.width = Pt(1.0)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.word_wrap = True
    return sh


def header(slide, s, style, idx, total, wide=CONTENT_W):
    """标题区：两版风格不同（B 用短下划线，C 用整条色带）。"""
    if style["header_band"] is not None:
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Inches(0.16),
                                      SLIDE_W, Inches(1.12))
        band.fill.solid(); band.fill.fore_color.rgb = style["header_band"]
        band.line.fill.background(); band.shadow.inherit = False
    if style["side_bar"] is not None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0),
                                     Inches(0.13), SLIDE_H)
        bar.fill.solid(); bar.fill.fore_color.rgb = style["side_bar"]
        bar.line.fill.background(); bar.shadow.inherit = False
    _, tf = tb(slide, MARGIN, TITLE_TOP, wide, Inches(0.8))
    para(tf, s.get("title", ""), size=style["title_size"], bold=True, color=style["ink"],
         first=True, before=0, spacing=1.05)
    if s.get("subtitle"):
        _, tf2 = tb(slide, MARGIN, SUB_TOP, wide, Inches(0.34))
        para(tf2, s["subtitle"], size=16, color=style["muted"], first=True, before=0)
    if style["header_band"] is None:
        hline(slide, MARGIN, SUB_TOP + Inches(0.34), Inches(1.5), style["accent"], Pt(2.2))
    return BODY_TOP


def takeaway(slide, text, style, top=TAKE_TOP, height=TAKE_H):
    if style["takeaway_style"] == "band":
        box = card(slide, MARGIN, top, CONTENT_W, height, style,
                   fill=RGBColor(0xEC, 0xF7, 0xF4), line=style["accent2"], radius=0.1)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.18)
        para(tf, text, size=17, bold=True, color=style["ink"], first=True, before=0, after=0,
             spacing=1.05)
    else:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top + Inches(0.10),
                                     Inches(0.055), Inches(0.36))
        bar.fill.solid(); bar.fill.fore_color.rgb = style["accent2"]
        bar.line.fill.background(); bar.shadow.inherit = False
        _, tf = tb(slide, MARGIN + Inches(0.16), top, CONTENT_W - Inches(0.16), height,
                   anchor=MSO_ANCHOR.MIDDLE)
        para(tf, text, size=17, bold=True, color=style["ink"], first=True, before=0, after=0,
             spacing=1.1)


def footer(slide, style, idx, total, footer_text):
    _, tf = tb(slide, SLIDE_W - Inches(1.65), FOOT_TOP, Inches(1.2), Inches(0.3))
    para(tf, f"{idx} / {total}", size=15, color=style["muted"], first=True,
         align=PP_ALIGN.RIGHT, before=0, after=0)
    _, tf2 = tb(slide, MARGIN, FOOT_TOP, Inches(9.0), Inches(0.3))
    para(tf2, footer_text, size=15, color=style["muted"], first=True, before=0, after=0)


def add_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def figure(slide, name, top, style, max_h=None, max_w=None):
    return mp.place_figure(slide, name, top=top, max_w=max_w or CONTENT_W,
                           max_h=max_h or (TAKE_TOP - Inches(0.1) - top))


# ----------------------------------------------------------------- 各页渲染
def render(slide, s, style, idx, total, meta):
    kind = s["type"]

    if kind == "title":
        if style["header_band"] is not None:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Inches(2.02),
                                          SLIDE_W, Inches(0.10))
            band.fill.solid(); band.fill.fore_color.rgb = style["accent"]
            band.line.fill.background(); band.shadow.inherit = False
        _, tf = tb(slide, MARGIN, Inches(1.10), CONTENT_W, Inches(0.5))
        para(tf, s.get("eyebrow", ""), size=19, bold=True, color=style["accent2"],
             first=True, align=PP_ALIGN.CENTER, before=0, after=0)
        _, tf = tb(slide, MARGIN, Inches(2.44), CONTENT_W, Inches(1.8))
        para(tf, s["title"], size=30, bold=True, color=style["ink"], first=True,
             align=PP_ALIGN.CENTER, before=0, after=8, spacing=1.15)
        para(tf, s.get("title2", ""), size=30, bold=True, color=style["ink"],
             align=PP_ALIGN.CENTER, before=0, after=0, spacing=1.15)
        _, tf = tb(slide, MARGIN, Inches(4.62), CONTENT_W, Inches(1.9))
        for line in (meta["presenter"], meta["advisor"], meta.get("major", ""),
                     f"汇报日期：{meta['date']}"):
            if not line:
                continue
            para(tf, line, size=17, color=style["ink"], align=PP_ALIGN.CENTER,
                 before=(0 if line is meta["presenter"] else 10), after=0)
        if style["side_bar"] is not None:
            hline(slide, MARGIN + Inches(3.5), Inches(4.42), CONTENT_W - Inches(7.0),
                  style["accent2"], Pt(2.2))
        return

    if kind == "end":
        _, tf = tb(slide, MARGIN, Inches(2.7), CONTENT_W, Inches(1.6))
        para(tf, s["title"], size=34, bold=True, color=style["ink"], first=True,
             align=PP_ALIGN.CENTER, before=0, after=14)
        para(tf, s.get("title2", ""), size=22, color=style["accent2"],
             align=PP_ALIGN.CENTER, before=0, after=0)
        _, tf = tb(slide, MARGIN, Inches(4.7), CONTENT_W, Inches(0.8))
        para(tf, f"汇报人：{meta['presenter'].split('　')[0]}　{meta['advisor'].split('　')[0]}",
             size=16, color=style["muted"], first=True, align=PP_ALIGN.CENTER,
             before=0, after=0)
        return

    top = header(slide, s, style, idx, total)

    if kind == "toc":
        _, tf = tb(slide, MARGIN, top - Inches(0.05), CONTENT_W, Inches(4.3))
        for k, item in enumerate(s.get("items", [])):
            rich(tf, [(f"{k + 1:02d}　", True, style["accent2"]), (item, False, style["ink"])],
                 size=style["toc_number_size"], first=(k == 0), before=(0 if k == 0 else 20))

    elif kind == "figure":
        fig_top = top - Inches(0.05)
        bottom = TAKE_TOP - Inches(0.95)
        figure(slide, s["figure"], fig_top, style, max_h=bottom - fig_top)
        if s.get("body"):
            _, tf = tb(slide, MARGIN, TAKE_TOP - Inches(0.92), CONTENT_W, Inches(0.85))
            para(tf, s["body"], size=16.5, color=style["ink"], first=True, before=0,
                 after=0, spacing=1.25)

    elif kind == "bullets":
        bullets = s.get("bullets", [])
        has_callout = bool(s.get("callout"))
        _, tf = tb(slide, MARGIN, top + Inches(0.05), CONTENT_W,
                   Inches(3.4 if has_callout else 4.3))
        for k, text in enumerate(bullets):
            tag, _, rest = text.partition("　")
            rich(tf, [(f"{tag}　" if rest else "", True, style["accent2"]),
                      (rest or tag, False, style["ink"])],
                 size=style["body_size"], first=(k == 0), before=(0 if k == 0 else 16))
        if has_callout:
            co = s["callout"]
            if style["takeaway_style"] == "band":
                box = card(slide, MARGIN, Inches(4.98), CONTENT_W, Inches(1.18), style,
                           fill=RGBColor(0xFD, 0xF2, 0xE3), line=ORANGE, radius=0.08)
            else:
                box = card(slide, MARGIN, Inches(4.98), CONTENT_W, Inches(1.18), style,
                           fill=None, line=style["accent"], radius=0.0)
            tf2 = box.text_frame
            tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf2.margin_left = tf2.margin_right = Inches(0.2)
            para(tf2, co["label"], size=17, bold=True, color=ORANGE, first=True,
                 before=0, after=4)
            para(tf2, co["text"], size=18, bold=True, color=style["ink"], before=0,
                 after=0, spacing=1.15)

    elif kind == "cards" and s.get("cards"):
        bullets, cards = s.get("bullets", []), s["cards"]
        _, tf = tb(slide, MARGIN, top + Inches(0.06), Inches(7.15),
                   TAKE_TOP - Inches(0.15) - (top + Inches(0.06)))
        for k, text in enumerate(bullets):
            rich(tf, [("▪　", False, style["accent2"]), (text, False, style["ink"])],
                 size=16.5, first=(k == 0), before=(0 if k == 0 else 15))
        y = top + Inches(0.06)
        card_h = Inches(1.02)
        for label, desc, tone in cards:
            col = style["accent2"] if tone == "B" else style["accent"]
            box = card(slide, MARGIN + Inches(7.45), y, Inches(4.64), card_h, style)
            box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            box.text_frame.margin_left = Inches(0.16)
            rich(box.text_frame, [(label + "　", True, col), (desc, False, style["ink"])],
                 size=18, first=True, before=0, spacing=1.05)
            y += card_h + Inches(0.16)

    elif kind == "cards" and s.get("cards_problems"):
        y = top + Inches(0.05)
        for tag, prob, fix in s["cards_problems"]:
            box = card(slide, MARGIN, y, CONTENT_W, Inches(1.28), style,
                       fill=(RGBColor(0xFB, 0xF6, 0xF0) if style["takeaway_style"] == "band"
                             else None),
                       line=(ORANGE if style["takeaway_style"] == "band" else style["accent"]),
                       radius=0.08)
            tf2 = box.text_frame
            tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf2.margin_left = tf2.margin_right = Inches(0.2)
            para(tf2, f"{tag}　{prob}", size=17, bold=True, color=style["ink"], first=True,
                 before=0, after=4, spacing=1.12)
            para(tf2, fix, size=16.5, color=style["ink"], before=0, after=0, spacing=1.12)
            y += Inches(1.42)

    elif kind == "steps":
        _, tf = tb(slide, MARGIN, top + Inches(0.1), CONTENT_W, Inches(4.2))
        for k, (tag, text) in enumerate(s.get("steps", [])):
            rich(tf, [(tag + "　", True, style["accent2"]), (text, False, style["ink"])],
                 size=style["body_size"], first=(k == 0), before=(0 if k == 0 else 18))

    if s.get("takeaway"):
        takeaway(slide, s["takeaway"], style)


def build(style_key: str, data: dict) -> Presentation:
    style = STYLES[style_key]
    meta = data["meta"]
    slides_data = data["slides"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    total = len(slides_data)
    for s in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        render(slide, s, style, s["n"], total, meta)
        if s["n"] not in (1, total):
            footer(slide, style, s["n"], total, meta["footer"])
        add_notes(slide, s.get("note"))
    return prs


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--style", choices=["B", "C", "both"], default="both")
    args = ap.parse_args()

    data = json.loads(OUTLINE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    keys = ["B", "C"] if args.style == "both" else [args.style]

    rc = 0
    for k in keys:
        prs = build(k, data)
        out = OUT_DIR / f"中期答辩_{k}_{STYLES[k]['name']}.pptx"
        prs.save(str(out))
        print(f"\n=== 风格 {k}（{STYLES[k]['name']}）-> {out}")
        rc |= mp.audit(out)
    print("\n说明：三版内容一致（同一份 docs/ppt_outline.json），仅版式与配色不同；"
          "均为原生可编辑 pptx，且每页写入了演讲备注。")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
