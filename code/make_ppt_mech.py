#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""重构机制补充页，解决重叠和排版丑的问题"""

from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_ppt_nature as N
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def clean_mech_append(prs, first_idx, total):
    """4页机制补充页，排版重构：每页文字精简，留白多，无重叠"""
    
    # M1: 七环逻辑链 - 图为主，文字精简为3行要点
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    N.title(slide, "抗菌肽与AD关联可拆为七环，证据强度分层", size=24)
    # 图占大部分
    N.picture(slide, "figM1_关联逻辑链.png", N.BODY_TOP, Inches(3.8))
    N.source_line(slide, "图：本项目自制，蓝=已有实验/临床证据，绿=计算可预测，橙=本课题待验证")
    # 底部要点精简 - 增加高度
    _, tf = N.tb(slide, N.LEFT, Inches(5.4), N.CONTENT_W, Inches(1.5))
    N.para(tf, "① AD与感染相关 ② 感染驱动抗菌肽/Aβ生成 ③ Aβ本身是抗菌肽 ④ 活性可计算预测 ⑤ AD组特异肽随阶段变化 ⑥ 双刃剑 ⑦ 正反馈环与AChE-PAS落点", size=15, color=N.BODY, first=True, spacing=1.25)
    N.footer(slide, first_idx, total)
    N.notes(slide, "七环逻辑链：把抗菌肽与AD的关联拆成七环，蓝色是已有实验或临床证据，绿色是计算方法能给出候选与优先序的部分，橙色是本课题要回答或需要验证的部分。老师问任何一环都能说清证据到哪一步。")

    # M2: 为什么AD更多 - 左右分栏，左侧3要点，右侧机制图示文字
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    N.title(slide, "为什么AD更多：炎症—Aβ—抗菌肽正反馈环", size=24)
    # 左侧要点 - 增加高度避免溢出
    _, tf = N.tb(slide, N.LEFT, Inches(1.6), Inches(7.5), Inches(4.0))
    N.para(tf, "上游驱动：菌群失衡与肠屏障↑，LPS持续入血", size=17, bold=True, color=N.INK, first=True, spacing=1.3)
    N.para(tf, "• 促炎菌 Escherichia/Shigella、Proteobacteria 富集", size=15, color=N.BODY, before=6, spacing=1.25)
    N.para(tf, "• 产丁酸菌 Eubacterium、Faecalibacterium 降低", size=15, color=N.BODY, before=4, spacing=1.25)
    N.para(tf, "• 血浆LPS约3倍，与单核活化正相关 [13,14]", size=15, color=N.BODY, before=4, spacing=1.25)
    N.para(tf, "中游放大：TLR4/NF-κB → 促炎因子 → 上调Aβ/抗菌肽", size=17, bold=True, color=N.INK, before=14, spacing=1.3)
    N.para(tf, "• LPS与Aβ均为TLR4/CD14激动剂", size=15, color=N.BODY, before=6, spacing=1.25)
    N.para(tf, "下游循环：Aβ/抗菌肽聚集 → 小胶质活化 → 更多炎症", size=17, bold=True, color=N.INK, before=14, spacing=1.3)
    N.para(tf, "• LL-37经CLIC1致ROS与神经毒性 [11]", size=15, color=N.BODY, before=6, spacing=1.25)
    N.para(tf, "• AChE-PAS促纤维化 [27-32]", size=15, color=N.BODY, before=4, spacing=1.25)

    # 右侧强调框
    N.rect(slide, Inches(8.6), Inches(1.6), Inches(4.0), Inches(4.0), N.PANEL, N.RULE)
    _, tf = N.tb(slide, Inches(8.8), Inches(1.7), Inches(3.6), Inches(3.8))
    N.para(tf, "抗菌保护假说", size=16, bold=True, color=N.RED, first=True, spacing=1.1)
    N.para(tf, "AD病理部分源于有益的先天免疫反应在长期失衡后被放大 [3]", size=15, color=N.BODY, before=8, spacing=1.25)
    N.para(tf, "双刃剑", size=16, bold=True, color=N.TEAL, before=14, spacing=1.1)
    N.para(tf, "抑制感染（生理）vs 促进聚集与炎症（病理）", size=15, color=N.BODY, before=8, spacing=1.25)
    N.para(tf, "口径", size=16, bold=True, color=N.TEAL, before=14, spacing=1.1)
    N.para(tf, "区分宿主/微生物源、局部/系统、阶段", size=15, color=N.MUTED, before=8, spacing=1.25)

    N.footer(slide, first_idx+1, total)
    N.notes(slide, "为什么AD更多：因为存在正反馈。菌群失衡与肠屏障通透性增加使LPS长期刺激宿主，经TLR4/NF-κB升高促炎因子，促炎因子又上调Aβ与抗菌肽生成；Aβ与抗菌肽聚集激活小胶质，产生更多炎症介质，循环加强。这是抗菌保护假说的现代表述。")

    # M3: AChE-Aβ机制落点 - 图+3假设卡片，无重叠
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    N.title(slide, "机制落点：AChE–Aβ复合物与三个可检验假设（多文献）", size=23)
    # 左图
    pic, _, _, _ = N.picture(slide, "figM2_AChE_Aβ_机制.png", N.BODY_TOP, Inches(3.6), left=N.LEFT, width=Inches(7.8))
    # 右侧3假设 - 全部≥15pt
    y0 = 1.45
    hyps = [
        ("H1 竞争PAS", "结合PAS，减少AChE诱导成核", "对接+MD；propidium竞争；酶活"),
        ("H2 结合Aβ", "结合Aβ，改变聚集路径", "ThT+电镜/AFM；MD残基；参照[10-12]"),
        ("H3 膜水平", "影响膜互作与小胶质识别", "脂膜MD+炎症因子；参照[11]"),
    ]
    for i,(t,d,v) in enumerate(hyps):
        top = y0 + i*1.55
        N.rect(slide, Inches(8.6), Inches(top), Inches(4.0), Inches(1.35), N.PANEL, N.TEAL if i<2 else N.RED)
        _, tf = N.tb(slide, Inches(8.8), Inches(top+0.05), Inches(3.6), Inches(1.25))
        N.para(tf, t, size=16, bold=True, color=N.INK, first=True, spacing=1.1)
        N.para(tf, d, size=15, color=N.BODY, before=4, spacing=1.2)
        N.para(tf, v, size=15, color=N.MUTED, before=4, spacing=1.2)

    # 底部文献链 - 精简且≥15pt
    _, tf = N.tb(slide, N.LEFT, Inches(5.3), N.CONTENT_W, Inches(0.6))
    N.para(tf, "文献链：Inestrosa 1996[27] → De Ferrari 2001[29] → Inestrosa 2008[30] → Atanasova 2020 MD[32]，344-361为第二界面", size=15, color=N.MUTED, first=True, spacing=1.2)

    N.footer(slide, first_idx+2, total)
    N.notes(slide, "机制落点不只一篇动力学：Inestrosa 1996年首次报道AChE加速聚集可被PAS配体propidium抑制，De Ferrari 2001鉴定疏水基序，Inestrosa 2008综述，Atanasova 2020做1μs MD显示344-361区段稳定且不受双位点抑制剂位阻，由此提出H1-H3。")

    # M4: 证据分级与抗菌活性预测补充
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    N.title(slide, "证据分级与抗菌活性预测的文献基础", size=24)
    
    # 左侧表格：证据分级精简版 - 全部≥15pt
    from pptx.util import Emu
    shape = slide.shapes.add_table(7, 3, N.LEFT, N.BODY_TOP+Inches(0.1), Inches(7.8), Inches(3.2))
    table = shape.table
    table.columns[0].width = Inches(3.4)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(2.6)
    headers = ["要点","强度","来源"]
    for c,t in enumerate(headers):
        cell = table.cell(0,c)
        cell.text = ""
        N.para(cell.text_frame, t, size=15, bold=True, color=N.INK, first=True, spacing=1.05)
        cell.fill.solid()
        cell.fill.fore_color.rgb = N.PANEL
    rows = [
        ("Aβ抗菌活性 AD脑>对照","强","Soscia 2010 [1]"),
        ("Aβ保护抗感染 寡聚化必需","强","Kumar 2016 [2]"),
        ("宿主AMP上调 hBD-1/LL-37","中-强","Williams 2013等"),
        ("AChE经PAS促聚集","强","Inestrosa 1996 [27]"),
        ("AChE-Aβ MD稳定 344-361","中","Atanasova 2020 [32]"),
        ("菌群失衡 促炎↑产丁酸↓","中","Vogt 2017等"),
    ]
    for r,(a,b,c) in enumerate(rows, start=1):
        for col,txt in enumerate((a,b,c)):
            cell = table.cell(r,col)
            cell.text = ""
            N.para(cell.text_frame, txt, size=15, color=N.BODY, first=True, spacing=1.05)

    # 右侧：抗菌活性预测 - 全部≥15pt
    N.rect(slide, Inches(8.6), N.BODY_TOP+Inches(0.1), Inches(4.0), Inches(3.2), N.PANEL, N.RULE)
    _, tf = N.tb(slide, Inches(8.8), Inches(1.5), Inches(3.6), Inches(3.0))
    N.para(tf, "抗菌活性预测为何可行", size=15, bold=True, color=N.TEAL, first=True, spacing=1.1)
    N.para(tf, "• Ma 2022 Nat Biotech：2349候选→216合成→181活性83% [22]", size=15, color=N.BODY, before=8, spacing=1.2)
    N.para(tf, "• Santos-Junior 2024 Cell：86万肽，100→79活性 [23]", size=15, color=N.BODY, before=6, spacing=1.2)
    N.para(tf, "• Wan 2024综述 [24]", size=15, color=N.BODY, before=6, spacing=1.2)
    N.para(tf, "本课题：三模型共识→去重→实验验证", size=15, bold=True, color=N.TEAL, before=12, spacing=1.1)
    N.para(tf, "纸片扩散初筛+MIC，设阳性/阴性对照", size=15, color=N.BODY, before=6, spacing=1.2)

    # 底部：谁更多 - 精简≥15pt
    _, tf = N.tb(slide, N.LEFT, Inches(5.0), N.CONTENT_W, Inches(1.4))
    N.para(tf, "正常人 vs AD谁更多：AD侧更高，区分宿主/微生物源、局部/系统、阶段；微生物源以数据为准。为什么多：感染—炎症—Aβ/抗菌肽正反馈，促炎菌富集驱动。", size=15, color=N.BODY, first=True, spacing=1.25)

    N.footer(slide, first_idx+3, total)
    N.notes(slide, "证据分级表说明每一环的强度，待验证一栏与本课题直接对应。抗菌活性预测部分补充文献：Ma 2022和Santos-Junior 2024的大规模验证，说明从序列预测活性是可行且被验证的路线。正常人vs AD谁更多：AD侧更高但需分层，微生物源以数据为准。")

    return first_idx+3

if __name__ == "__main__":
    # test
    import json
    outline = json.loads((Path(__file__).parent.parent / "docs" / "ppt_outline.json").read_text(encoding="utf-8"))
    prs = N.build(outline, total=20)
    clean_mech_append(prs, first_idx=17, total=20)
    out = Path(__file__).parent.parent / "deliverable" / "中期答辩_H_nature风.pptx"
    prs.save(str(out))
    print(f"wrote {out} slides={len(prs.slides)}")


def append(prs, first_idx, total):
    return clean_mech_append(prs, first_idx, total)
