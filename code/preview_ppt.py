#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""preview_ppt.py - render an approximate PNG preview of every slide.

This environment has no PowerPoint and no LibreOffice, so the deck cannot be
rendered exactly. This script draws the shapes with PIL using the same
coordinates, colours and the real CJK font, which is accurate enough to spot
overlaps, text that runs out of its box, and missing images.

    python code/preview_ppt.py deliverable/中期答辩_A_学术蓝.pptx -o build/ppt_preview
"""

from __future__ import annotations

import argparse
import glob
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

DPI = 110


def font_for(size_pt: float, bold: bool = False):
    cands = sorted(glob.glob("/tmp/fonts/*CJK*.otf")) + \
            sorted(glob.glob("/tmp/fonts/*.otf"))
    px = max(8, int(round(size_pt * DPI / 72)))
    if cands:
        return ImageFont.truetype(cands[0], px)
    return ImageFont.load_default()


def page_bg(slide):
    """Read the slide's own <p:bg> fill, falling back to layout and master.

    ppt-master's native export promotes a full-canvas rect to the slide
    background, so a dark deck would otherwise preview as white.
    """
    from pptx.oxml.ns import qn

    def srgb(el):
        if el is None:
            return None
        clr = el.find(".//" + qn("a:srgbClr"))
        return clr.get("val") if clr is not None else None

    for node in (slide._element, slide.slide_layout._element,
                 slide.slide_layout.slide_master._element):
        bg = node.find(".//" + qn("p:bg"))
        v = srgb(bg)
        if v:
            return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
    return None


def pt2px(v) -> float:
    """Length (EMU) -> pixels at DPI"""
    if v is None:
        return 0.0
    return float(v.pt) * DPI / 72.0


def draw_text_frame(draw: ImageDraw.ImageDraw, shape, scale_x, scale_y):
    tf = shape.text_frame
    x0 = pt2px(shape.left) * scale_x
    y0 = pt2px(shape.top) * scale_y
    w = pt2px(shape.width) * scale_x
    ml = pt2px(getattr(tf, "margin_left", None))
    y = y0 + pt2px(getattr(tf, "margin_top", None))
    # vertical centring inside autoshapes
    total_h = 0
    prepared = []
    for p in tf.paragraphs:
        runs = [(r.text, (r.font.size.pt if r.font.size else 18),
                 bool(r.font.bold), r.font.color.rgb if r.font.color and r.font.color.type is not None else None)
                for r in p.runs]
        if not runs:
            prepared.append(None)
            continue
        size = max(r[1] for r in runs)
        text = "".join(r[0] for r in runs)
        font = font_for(size)
        # wrap
        lines, cur = [], ""
        for ch in text:
            if draw.textlength(cur + ch, font=font) > w - ml * 2 and cur:
                lines.append(cur)
                cur = ch
            else:
                cur += ch
        lines.append(cur)
        spacing = p.line_spacing or 1.2
        lh = size * spacing * 1.22 * DPI / 72
        prepared.append((lines, size, lh, runs[0][3], bool(runs[0][2]),
                         (p.space_before.pt if p.space_before else 0) * DPI / 72))
        total_h += len(lines) * lh + (p.space_before.pt if p.space_before else 0) * DPI / 72
    if total_h and shape.shape_type not in (None,) and shape.has_text_frame and shape.fill.type is not None:
        box_h = pt2px(shape.height) * scale_y
        if total_h < box_h:                      # middle anchor used by our bars
            y = y0 + (box_h - total_h) / 2
    for item in prepared:
        if item is None:
            continue
        lines, size, lh, color, bold, sb = item
        y += sb
        f = font_for(size, bold)
        for line in lines:
            draw.text((x0 + ml, y), line, font=f,
                      fill=tuple(color) if color else (17, 50, 79))
            y += lh


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", nargs="?", default="deliverable/中期答辩_A_学术蓝.pptx")
    ap.add_argument("-o", "--out", default="build/ppt_preview")
    args = ap.parse_args(argv)

    prs = Presentation(args.pptx)
    W = int(pt2px(prs.slide_width))
    H = int(pt2px(prs.slide_height))
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), page_bg(slide) or "white")
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                import io
                pic = Image.open(io.BytesIO(sh.image.blob))
                pic = pic.convert("RGB").resize(
                    (max(1, int(pt2px(sh.width))), max(1, int(pt2px(sh.height)))))
                img.paste(pic, (int(pt2px(sh.left)), int(pt2px(sh.top))))
                continue
            try:
                if sh.fill.type is not None and sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                    rgb = sh.fill.fore_color.rgb
                    x0, y0 = pt2px(sh.left), pt2px(sh.top)
                    x1, y1 = x0 + pt2px(sh.width), y0 + pt2px(sh.height)
                    if x1 - x0 > 1 and y1 - y0 > 1:
                        if "ROUNDED" in str(sh.shape_type):
                            d.rounded_rectangle([x0, y0, x1, y1], radius=9,
                                                fill=tuple(rgb), outline=(47, 111, 176), width=2)
                        else:
                            d.rectangle([x0, y0, x1, y1], fill=tuple(rgb))
            except Exception:
                pass
            if sh.has_text_frame:
                draw_text_frame(d, sh, 1.0, 1.0)
        img.save(out / f"slide{idx:02d}.png")
    print(f"{len(prs.slides._sldIdLst)} preview images -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
