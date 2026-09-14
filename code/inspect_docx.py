#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""inspect_docx.py - dump the structure of a .docx so we can fill it in later
without touching its original formatting.

Usage
-----
    python code/inspect_docx.py sources/中期.docx
    python code/inspect_docx.py sources/中期.docx -o docs/_template
    python code/inspect_docx.py sources/中期.docx --json-only

Output
------
    <out>/<name>.txt   human-readable outline (paragraph index / style / text,
                       table size and every cell)
    <out>/<name>.json  machine-readable structure (paragraphs, runs, fonts,
                       sizes, tables, sections, headers/footers)

Every body element gets a stable index (paragraphs and tables share one
sequence), which is what fill_docx.py uses as an anchor.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# runs that look like "please fill here" placeholders
PLACEHOLDER_RE = re.compile(r"[_＿﹍]{2,}|[．.]{4,}|…{2,}|【[^】]*】|\(\s*请|（\s*请")
BLANK_RE = re.compile(r"[_＿﹍]{2,}|[．.]{4,}|…{2,}")


def iter_body(doc):
    """Yield ('p', Paragraph) / ('tbl', Table) in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield "tbl", Table(child, doc)


def run_info(run) -> dict:
    f = run.font
    info = {"text": run.text}
    if f.name:
        info["font"] = f.name
    # east-asian font lives in the rPr, python-docx does not expose it directly
    rPr = run._element.find(qn("w:rPr"))
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            ea = rFonts.get(qn("w:eastAsia"))
            if ea:
                info["font_eastasia"] = ea
    if f.size is not None:
        info["size_pt"] = f.size.pt
    if f.bold:
        info["bold"] = True
    if f.italic:
        info["italic"] = True
    if f.underline:
        info["underline"] = True
    try:
        if f.color is not None and f.color.rgb is not None:
            info["color"] = str(f.color.rgb)
    except Exception:
        pass
    return info


def para_info(p: Paragraph, idx: int) -> dict:
    d = {
        "index": idx,
        "style": p.style.name if p.style is not None else None,
        "text": p.text,
        "runs": [run_info(r) for r in p.runs],
    }
    if p.alignment is not None:
        d["align"] = str(p.alignment)
    pf = p.paragraph_format
    if pf.first_line_indent is not None:
        d["first_line_indent_pt"] = pf.first_line_indent.pt
    if pf.left_indent is not None:
        d["left_indent_pt"] = pf.left_indent.pt
    if pf.line_spacing is not None:
        d["line_spacing"] = pf.line_spacing
    d["is_placeholder"] = bool(PLACEHOLDER_RE.search(p.text))
    d["num_runs"] = len(p.runs)
    return d


def cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def table_info(t: Table, idx: int) -> dict:
    rows = []
    for r_i, row in enumerate(t.rows):
        cells = []
        for c_i, cell in enumerate(row.cells):
            cells.append(
                {
                    "row": r_i,
                    "col": c_i,
                    "text": cell_text(cell),
                    "paragraphs": [
                        {
                            "text": p.text,
                            "style": p.style.name if p.style is not None else None,
                            "runs": [run_info(r) for r in p.runs],
                        }
                        for p in cell.paragraphs
                    ],
                }
            )
        rows.append(cells)
    return {
        "index": idx,
        "type": "table",
        "n_rows": len(t.rows),
        "n_cols": len(t.columns),
        "rows": rows,
    }


def section_info(doc) -> list:
    out = []
    for s_i, s in enumerate(doc.sections):
        d = {"section": s_i, "orientation": str(s.orientation), "page_w_pt": s.page_width.pt if s.page_width else None,
             "page_h_pt": s.page_height.pt if s.page_height else None}
        d["header"] = [p.text for p in s.header.paragraphs]
        d["footer"] = [p.text for p in s.footer.paragraphs]
        out.append(d)
    return out


def build(docx_path: Path) -> dict:
    doc = Document(str(docx_path))
    elements = []
    i = 0
    for kind, obj in iter_body(doc):
        if kind == "p":
            elements.append(para_info(obj, i))
        else:
            elements.append(table_info(obj, i))
        i += 1

    return {
        "file": str(docx_path),
        "n_body_elements": len(elements),
        "n_paragraphs": sum(1 for e in elements if e.get("type") != "table"),
        "n_tables": sum(1 for e in elements if e.get("type") == "table"),
        "sections": section_info(doc),
        "elements": elements,
    }


def to_text(struct: dict) -> str:
    lines = []
    lines.append(f"# {struct['file']}")
    lines.append(f"# body elements: {struct['n_body_elements']} "
                 f"(paragraphs {struct['n_paragraphs']}, tables {struct['n_tables']})")
    lines.append("")
    for e in struct["elements"]:
        if e.get("type") == "table":
            lines.append(f"[tbl {e['index']}] table {e['n_rows']}x{e['n_cols']}")
            for r in e["rows"]:
                for c in r:
                    txt = c["text"].replace("\n", " ⏎ ")
                    lines.append(f"    (r{c['row']},c{c['col']}) {txt}")
            lines.append("")
            continue
        style = e.get("style") or "-"
        align = e.get("align") or "-"
        size = next((f" {r['size_pt']}pt" for r in e["runs"] if "size_pt" in r), "")
        font = next((f" {r.get('font') or r.get('font_eastasia')}" for r in e["runs"]
                     if r.get("font") or r.get("font_eastasia")), "")
        mark = "  <== PLACEHOLDER" if e.get("is_placeholder") else ""
        txt = e["text"].replace("\n", " ⏎ ")
        lines.append(f"[p {e['index']:>3}] {style:<12} {align:<10}{font}{size} | {txt}{mark}")
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="dump .docx structure as text + json")
    ap.add_argument("docx", help="path to the .docx file")
    ap.add_argument("-o", "--out-dir", default="docs/_template", help="output directory")
    ap.add_argument("--json-only", action="store_true")
    args = ap.parse_args(argv)

    src = Path(args.docx)
    if not src.exists():
        print(f"[ERROR] not found: {src}", file=sys.stderr)
        return 1

    struct = build(src)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = src.stem

    json_path = out_dir / f"{stem}.json"
    json_path.write_text(json.dumps(struct, ensure_ascii=False, indent=1), encoding="utf-8")

    if not args.json_only:
        txt_path = out_dir / f"{stem}.txt"
        txt_path.write_text(to_text(struct), encoding="utf-8")
        print(f"wrote {txt_path}")

    print(f"wrote {json_path}")
    print(f"paragraphs {struct['n_paragraphs']}, tables {struct['n_tables']}, "
          f"placeholders {sum(1 for e in struct['elements'] if e.get('is_placeholder'))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
