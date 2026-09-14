#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_ppt.py - build the 中期答辩 PPT (defence deck).

Design rules
------------
* 16:9, conclusion-style titles, one idea per slide
* Chinese is native PowerPoint text (font 微软雅黑, east-asian typeface set
  explicitly) so it renders correctly on any Windows machine
* NOTHING on a slide is smaller than 15 pt - including the text inside the
  figures (make_figures.py guarantees >= 15 pt effective when scaled to the
  slide width, this script re-checks every native run)
* layout is fixed by constants below (TITLE_TOP / CONTENT_TOP / TAKEAWAY_TOP /
  FOOTER_TOP) so pictures, conclusion bars and footers can never overlap

    python code/make_ppt.py            # -> deliverable/中期答辩.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ------------------------------------------------------------------ geometry
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_W = SLIDE_W - MARGIN * 2          # 12.09 in
TITLE_TOP = Inches(0.32)
SUBTITLE_TOP = Inches(1.10)
CONTENT_TOP = Inches(1.62)
TAKEAWAY_TOP = Inches(6.30)
TAKEAWAY_H = Inches(0.60)
FOOTER_TOP = Inches(7.04)
CONTENT_BOTTOM = TAKEAWAY_TOP - Inches(0.10)

MIN_PT = 15

# ------------------------------------------------------------------ colours
INK = RGBColor(0x12, 0x32, 0x4F)
BLUE = RGBColor(0x2F, 0x6F, 0xB0)
TEAL = RGBColor(0x2F, 0x9E, 0x8F)
ORANGE = RGBColor(0xE0, 0x8A, 0x2E)
GREY = RGBColor(0x6B, 0x7C, 0x8C)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MINT = RGBColor(0xEC, 0xF7, 0xF4)
SAND = RGBColor(0xFD, 0xF2, 0xE3)

CN_FONT = "微软雅黑"
EN_FONT = "Arial"
FIG = Path("results/figures")
OUT = Path("deliverable/中期答辩.pptx")

TITLE_MAX = 32
PROBLEMS: list[str] = []


# ------------------------------------------------------------------ helpers
def set_run_font(run, size: float, bold=False, color: RGBColor = INK,
                 cn: str = CN_FONT, en: str = EN_FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = en
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", cn), ("a:cs", en)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def add_para(tf, text, size=17, bold=False, color=INK, space_before=6, space_after=2,
             align=PP_ALIGN.LEFT, first=False, line_spacing=1.18):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    set_run_font(r, size, bold, color)
    return p


def add_rich(tf, parts, size=17, space_before=6, space_after=2, first=False,
             line_spacing=1.2, align=PP_ALIGN.LEFT):
    """parts = [(text, bold, colour), ...]"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for text, bold, color in parts:
        r = p.add_run()
        r.text = text
        set_run_font(r, size, bold, color)
    return p


def add_title(slide, text, sub=None):
    if len(text) > TITLE_MAX:
        PROBLEMS.append(f"title too long ({len(text)}): {text}")
    tb, tf = textbox(slide, MARGIN, TITLE_TOP, CONTENT_W, Inches(0.72))
    add_para(tf, text, size=27, bold=True, color=INK, first=True, space_before=0,
             line_spacing=1.05)
    if sub:
        tb2, tf2 = textbox(slide, MARGIN, SUBTITLE_TOP, CONTENT_W, Inches(0.36))
        add_para(tf2, sub, size=16, color=GREY, first=True, space_before=0)
    return CONTENT_TOP


def add_round_box(slide, left, top, width, height, fill=LIGHT, line=BLUE, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.word_wrap = True
    return sh


def add_takeaway(slide, text):
    box = add_round_box(slide, MARGIN, TAKEAWAY_TOP, CONTENT_W, TAKEAWAY_H, fill=MINT, line=TEAL)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.18)
    add_para(tf, text, size=17, bold=True, color=INK, first=True, space_before=0, space_after=0,
             line_spacing=1.05)
    return box


def add_footer(slide, idx, total):
    tb, tf = textbox(slide, SLIDE_W - Inches(1.65), FOOTER_TOP, Inches(1.2), Inches(0.3))
    add_para(tf, f"{idx} / {total}", size=15, color=GREY, first=True,
             align=PP_ALIGN.RIGHT, space_before=0, space_after=0)
    tb2, tf2 = textbox(slide, MARGIN, FOOTER_TOP, Inches(8.6), Inches(0.3))
    add_para(tf2, "研究生论文中期检查 · 基于深度学习的 AD 肠道微生物组抗菌肽差异研究",
             size=15, color=GREY, first=True, space_before=0, space_after=0)


def place_figure(slide, name, top=CONTENT_TOP, max_w=CONTENT_W, max_h=None, center=True,
                 left=None):
    from PIL import Image
    path = FIG / name
    with Image.open(path) as im:
        w, h = im.size
    avail_h = (max_h or (CONTENT_BOTTOM - top))
    scale = min(max_w / w, avail_h / h)
    dw, dh = int(w * scale), int(h * scale)
    if left is not None:
        l = int(left)
    else:
        l = int(MARGIN + (CONTENT_W - dw) / 2) if center else int(MARGIN)
    slide.shapes.add_picture(str(path), l, top, dw, dh)
    return top + dh


def number_cards(slide, cards, left, top, width, height):
    """cards = [(big, label, colour), ...] stacked vertically"""
    n = len(cards)
    gap = Inches(0.16)
    card_h = int((height - gap * (n - 1)) / n)
    y = top
    for big, label, color in cards:
        box = add_round_box(slide, left, y, width, card_h, fill=LIGHT, line=color, radius=0.12)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.16)
        add_rich(tf, [(big + "　", True, color), (label, False, INK)], size=18,
                 first=True, space_before=0, space_after=0, line_spacing=1.05)
        y += card_h + gap


def bullet_block(slide, bullets, left, top, width, size=16.5, gap=12):
    tb, tf = textbox(slide, left, top, width, CONTENT_BOTTOM - top)
    for k, text in enumerate(bullets):
        add_rich(tf, [("▪　", False, BLUE), (text, False, INK)], size=size,
                 first=(k == 0), space_before=(0 if k == 0 else gap), space_after=0,
                 line_spacing=1.22)
    return tb


# ------------------------------------------------------------------ deck
def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slides = []

    # 1 ─ 封面 -------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN + Inches(3.5), Inches(2.16),
                            CONTENT_W - Inches(7.0), Pt(2.6))
    ln.fill.solid(); ln.fill.fore_color.rgb = BLUE
    ln.line.fill.background(); ln.shadow.inherit = False
    tb, tf = textbox(s, MARGIN, Inches(1.06), CONTENT_W, Inches(0.5))
    add_para(tf, "研究生论文中期检查", size=19, bold=True, color=BLUE, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0)
    tb, tf = textbox(s, MARGIN, Inches(2.46), CONTENT_W, Inches(1.7))
    add_para(tf, "基于深度学习的阿尔茨海默症患者与健康人群", size=30, bold=True,
             color=INK, first=True, align=PP_ALIGN.CENTER, space_before=0, space_after=8,
             line_spacing=1.15)
    add_para(tf, "肠道微生物组中抗菌肽的差异性研究", size=30, bold=True, color=INK,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0, line_spacing=1.15)
    tb, tf = textbox(s, MARGIN, Inches(4.7), CONTENT_W, Inches(1.8))
    for line in ["汇报人：【姓名】　　学号：【学号】",
                 "指导教师：【导师姓名】　　培养单位：【学院】",
                 "汇报日期：2026 年 9 月"]:
        add_para(tf, line, size=17, color=INK, align=PP_ALIGN.CENTER, space_before=10,
                 space_after=0)
    slides.append(s)

    # 2 ─ 提纲 -------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "汇报提纲")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP - Inches(0.05), CONTENT_W, Inches(4.3))
    items = [
        ("01", "研究背景与科学问题：AD—肠—脑轴—抗菌肽"),
        ("02", "研究目标与技术路线"),
        ("03", "已完成工作：476 例宏基因组 → 9 139.2 万条非冗余多肽"),
        ("04", "阶段性结果：四个队列的 sORF 数量层面比较"),
        ("05", "下一步计划：DeepMetaAMP 模型与候选抗菌肽差异分析"),
        ("06", "研究进度、存在问题与应对措施"),
    ]
    for k, (num, text) in enumerate(items):
        add_rich(tf, [(num + "　", True, TEAL), (text, False, INK)], size=20,
                 first=(k == 0), space_before=(0 if k == 0 else 20), space_after=0)
    add_takeaway(s, "汇报重点：已完成的数据资源与阶段性结果，以及下一步的技术方案")
    slides.append(s)

    # 3 ─ 背景 -------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "AD 与肠道微生物组通过“肠—脑轴”双向关联", "研究背景与科学问题")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP, CONTENT_W, Inches(3.3))
    rows = [
        ("现状 1", "AD 患者肠道微生物组存在菌群失调：短链脂肪酸产生菌减少、炎症相关菌增多，"
                   "伴随肠屏障功能下降与系统性炎症升高。"),
        ("现状 2", "抗菌肽（AMPs）兼具直接抗菌与免疫调节功能，是连接微生物刺激与神经炎症的"
                   "潜在效应分子；微生物基因组中的小开放阅读框（sORF）可编码此类短肽。"),
        ("现状 3", "现有 AMPs 研究多依赖宿主来源数据库与通用预测模型，缺乏面向 AD 患者、"
                   "覆盖病程各阶段、且基于微生物组序列空间的疾病特异性系统分析。"),
    ]
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "　", True, BLUE), (text, False, INK)], size=17,
                 first=(k == 0), space_before=(0 if k == 0 else 15), space_after=0,
                 line_spacing=1.25)
    box = add_round_box(s, MARGIN, Inches(5.0), CONTENT_W, Inches(1.16), fill=SAND, line=ORANGE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    add_para(tf, "科学问题", size=17, bold=True, color=ORANGE, first=True,
             space_before=0, space_after=4)
    add_para(tf, "AD 患者与健康人群肠道微生物组所编码的候选抗菌肽是否存在差异？"
                 "若存在，是否随认知功能下降的病程阶段呈现连续演变？",
             size=18, bold=True, color=INK, space_before=0, space_after=0, line_spacing=1.15)
    slides.append(s)

    # 4 ─ 目标与技术路线 ---------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "研究目标：构建无偏多肽库并识别 AD 相关候选抗菌肽", "研究目标与技术路线")
    place_figure(s, "fig1_技术路线.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "技术路线两端已打通：左侧数据处理与多肽库构建已完成，右侧模型与差异分析正在进行")
    slides.append(s)

    # 5 ─ 已完成 ① ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成 ①：22 582 个 MAGs，1 971 个代表基因组", "已完成工作 · 数据资源")
    bullets = [
        "476 例粪便双端鸟枪法数据：fastp 质控 + KneadData/Bowtie2 去宿主（hg38）→ MEGAHIT 从头组装",
        "MetaBAT2 / MaxBin2 / CONCOCT 三算法独立分箱，metaWRAP bin_refinement 提纯（完整度 ≥50%、污染度 ≤10%）",
        "CheckM 谱系标记基因 + barrnap + tRNAscan-SE 按 MIMAG 标准分级，757 个达到高质量标准",
        "dRep 在 95% ANI 阈值下去冗余得到 1 971 个种水平代表基因组，覆盖 459/476 例样本（96.4%）",
        "参考集不用 757 个高质量基因组：那样去重后仅剩 80 个物种、98 例样本无基因组入选",
    ]
    bullet_block(s, bullets, MARGIN, CONTENT_TOP + Inches(0.12), Inches(7.15))
    number_cards(s, [
        ("476", "例粪便宏基因组样本", BLUE),
        ("22 582", "个组装基因组（MAGs）", BLUE),
        ("757", "个 MIMAG 高质量基因组", TEAL),
        ("1 971", "个种水平代表基因组", TEAL),
    ], left=MARGIN + Inches(7.45), top=CONTENT_TOP + Inches(0.12),
        width=Inches(4.64), height=Inches(4.4))
    add_takeaway(s, "参考集以物种多样性为准：1 971 个代表基因组对应 459 例样本，保证下游分析无系统性偏倚")
    slides.append(s)

    # 6 ─ 已完成 ② ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成 ②：9 139.2 万条非冗余微生物源多肽库", "已完成工作 · 多肽库构建")
    rows = [
        ("预测工具", "EMBOSS getorf v6.6.0.0（-find 0 -table 11 -minsize 15 -maxsize 150）"),
        ("长度说明", "长度参数以核苷酸计，15—150 nt 即 5—50 aa"),
        ("产出规模", "2.19 亿条原始 sORF → 精确去冗余后 9 139.2 万条"),
        ("无偏性", "未引入抗菌肽数据库过滤或分类器，序列空间无偏"),
        ("可溯源性", "序列头含 MAG 标识，可回溯至来源代表基因组与样本"),
        ("基因组间差异", "1.99 万—42.5 万条/基因组，极差约 21.3 倍"),
    ]
    tb, tf = textbox(s, MARGIN, CONTENT_TOP + Inches(0.06), Inches(6.95),
                     CONTENT_BOTTOM - CONTENT_TOP)
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "：", True, TEAL), (text, False, INK)], size=16.5,
                 first=(k == 0), space_before=(0 if k == 0 else 16), space_after=0,
                 line_spacing=1.2)
    fig_top = place_figure(s, "fig2_产出规模.png", top=CONTENT_TOP + Inches(0.05),
                           max_w=Inches(4.55), max_h=Inches(2.5), center=False,
                           left=Inches(8.0))
    tb, tf = textbox(s, Inches(8.0), CONTENT_TOP + Inches(2.55), Inches(4.7), Inches(1.5))
    add_para(tf, "该多肽库覆盖队列中全部 459 例有效样本，"
                 "是后续抗菌肽预测与组间差异分析的统一输入。",
             size=16, color=GREY, first=True, space_before=0, space_after=0, line_spacing=1.25)
    add_takeaway(s, "以无偏方式构建的多肽库是本研究的核心数据资源，规模与可溯源性均满足下游模型需求")
    slides.append(s)

    # 7 ─ 已完成 ③ ---------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "已完成 ③：严格匹配亚队列与四个平行分析队列", "已完成工作 · 队列设计")
    place_figure(s, "fig3_队列设计.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "匹配亚队列各阶段均 53 例、性别构成一致、平均年龄最大差 1.8 岁，可分离病程效应与增龄效应")
    slides.append(s)

    # 8 ─ 结果 ① -----------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "阶段性结果 ①：sORF 数量指标无统计学差异", "阶段性结果 · 数量层面")
    place_figure(s, "fig4_sORF数量.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "主要结局 sORF_per_MAG 的 P 值为 0.7344（匹配队列）与 0.8359（全队列），组间差异均不显著")
    slides.append(s)

    # 9 ─ 结果 ② -----------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "阶段性结果 ②：12 项检验均未达显著水平", "阶段性结果 · 检验汇总")
    place_figure(s, "fig5_P值.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "五个阶段与 NC-AD 二元比较在匹配队列和全队列中结论一致，提示结果稳健")
    slides.append(s)

    # 10 ─ 结果解读 --------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "结果解读：组间波动源于组装质量差异", "阶段性结果 · 解读与启示")
    tb, tf = textbox(s, MARGIN, CONTENT_TOP, CONTENT_W, Inches(3.6))
    rows = [
        ("波动来源", "Total_sORF 的组间波动与代表基因组数 N_MAG 同向变化（匹配队列中 MCI 均值 3.5、"
                     "AD 均值 4.7），提示由贡献基因组数量驱动。"),
        ("方向不一致", "sORF_per_MAG 在匹配队列中 AD 略高于 NC（114 555.0 vs 113 306.7），"
                       "在全队列中则相反（111 582.5 vs 114 529.3），支持组间无真实差异。"),
        ("效应量对比", "组间相对差异为 9.4%（匹配队列）与 5.5%（全队列），"
                       "远小于单个基因组间 sORF 产出量的固有差异（约 21.3 倍）。"),
        ("研究启示", "AD 病程中微生物源 sORF 的总体编码容量保持稳定；差异若存在，"
                     "更可能体现于序列组成与功能亚类，因此后续分析须转向候选抗菌肽层面。"),
    ]
    for k, (tag, text) in enumerate(rows):
        add_rich(tf, [(tag + "：", True, BLUE), (text, False, INK)], size=17,
                 first=(k == 0), space_before=(0 if k == 0 else 15), space_after=0,
                 line_spacing=1.25)
    add_takeaway(s, "阴性结果明确了技术方向：下一步聚焦候选抗菌肽的组成差异与阶段演变")
    slides.append(s)

    # 11 ─ 模型 ------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "下一步 ①：构建 DeepMetaAMP 抗菌肽预测模型", "下一步计划 · 模型构建")
    place_figure(s, "fig7_模型结构.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "模型以 ESM-2 序列表征为核心，结合注意力池化与 Focal Loss 应对正负样本不平衡")
    slides.append(s)

    # 12 ─ 分析方案 --------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "下一步 ②：全量推理与四队列差异分析方案", "下一步计划 · 分析方案")
    steps = [
        ("1. 批量推理", "对 9 139.2 万条多肽分批预测，按验证集阈值筛选高置信度候选；与 DRAMP / APD3 / "
                        "AMPSphere 比对剔除已知家族，保留新颖候选。"),
        ("2. 丰度定量", "以 Bowtie2/BWA 将各样本 clean reads 回贴至 1 971 个代表基因组，CoverM 计算候选 "
                        "sORF 覆盖度与丰度，构建样本 × 候选多肽矩阵并做 CLR 变换。"),
        ("3. 组间差异分析", "四个队列分别用 Kruskal-Wallis H（五阶段）与 Mann-Whitney U（NC-AD）检验，"
                            "以 BH 法控制错误发现率（FDR），以 Cliff's delta 报告效应量；PCoA 与 PERMANOVA "
                            "评估整体组成差异，随机森林评估判别能力。"),
        ("4. 趋势与可解释性", "Jonckheere-Terpstra 检验评估随病程的单调趋势；与 MMSE、MoCA 评分做偏相关"
                              "（校正年龄、性别、BMI 与用药）；注意力权重与 SHAP 定位关键位点，MEME 挖掘保守基序。"),
    ]
    tb, tf = textbox(s, MARGIN, CONTENT_TOP, CONTENT_W, Inches(4.2))
    for k, (tag, text) in enumerate(steps):
        add_rich(tf, [(tag + "　", True, ORANGE), (text, False, INK)], size=16.5,
                 first=(k == 0), space_before=(0 if k == 0 else 15), space_after=0,
                 line_spacing=1.22)
    add_takeaway(s, "沿用已完成的数据资源与统计框架，无需重复数据生产，剩余工作量可控")
    slides.append(s)

    # 13 ─ 进度 ------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "研究进度：整体约 65%，与开题计划一致", "研究进度与工作量")
    place_figure(s, "fig6_进度甘特.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "前 5 项任务全部完成（100%）；模型训练已完成数据集整理与框架搭建，完成度约 35%")
    slides.append(s)

    # 14 ─ 下一步计划 ------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "剩余约 35% 工作的安排与交付物", "下一步计划")
    place_figure(s, "fig8_下一步计划.png", max_h=CONTENT_BOTTOM - CONTENT_TOP)
    add_takeaway(s, "论文撰写与预答辩安排在最后阶段，计划 2027 年 1—3 月完成学位论文初稿")
    slides.append(s)

    # 15 ─ 问题与对策 ------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "存在问题与应对措施", "问题与对策")
    cards = [
        ("问题 1", "去冗余后 17 例受试者无代表基因组入选（覆盖 459/476 例）",
         "应对：丰度层面改用全量 22 582 个 MAGs 回贴，恢复样本覆盖，避免参考集筛选造成样本缺失。"),
        ("问题 2", "候选抗菌肽存在假阳性风险",
         "应对：同源过滤 + 理化性质一致性评估 + 外部独立测试集三重策略控制，并尽可能辅以体外抑菌实验验证。"),
        ("问题 3", "深度学习模型在宏基因组长尾序列上的泛化能力",
         "应对：多次随机种子重复训练与独立数据集验证，报告性能的不确定性区间；必要时以集成策略提升稳健性。"),
    ]
    y = CONTENT_TOP + Inches(0.05)
    for tag, prob, fix in cards:
        box = add_round_box(s, MARGIN, y, CONTENT_W, Inches(1.28), fill=SAND, line=ORANGE)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.2)
        add_para(tf, f"{tag}　{prob}", size=17, bold=True, color=INK, first=True,
                 space_before=0, space_after=4, line_spacing=1.12)
        add_para(tf, fix, size=16.5, color=INK, space_before=0, space_after=0, line_spacing=1.12)
        y += Inches(1.42)
    add_takeaway(s, "问题均有明确的技术对策，可通过在研工作按期解决")
    slides.append(s)

    # 16 ─ 致谢 ------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb, tf = textbox(s, MARGIN, Inches(2.6), CONTENT_W, Inches(1.5))
    add_para(tf, "恳请各位老师批评指正", size=34, bold=True, color=INK, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=14)
    add_para(tf, "谢谢！", size=22, color=BLUE, align=PP_ALIGN.CENTER, space_before=0,
             space_after=0)
    tb, tf = textbox(s, MARGIN, Inches(4.6), CONTENT_W, Inches(0.8))
    add_para(tf, "汇报人：【姓名】　指导教师：【导师姓名】", size=16, color=GREY, first=True,
             align=PP_ALIGN.CENTER, space_before=0, space_after=0)
    slides.append(s)

    for i, sl in enumerate(slides, 1):
        if i not in (1, len(slides)):
            add_footer(sl, i, len(slides))
    return prs


# ------------------------------------------------------------------ checks
def _cjk_font(size_pt: float):
    import glob as _glob
    from PIL import ImageFont
    cands = sorted(_glob.glob("/tmp/fonts/*CJK*.otf"))
    if not cands:
        return None
    return ImageFont.truetype(cands[0], max(8, int(round(size_pt * 96 / 72))))


def est_text_height_in(sp, box_w_in: float) -> float:
    """Estimated height (inches) of the wrapped text of a shape."""
    if not sp.has_text_frame:
        return 0.0
    total_pt = 0.0
    for p in sp.text_frame.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = max((r.font.size.pt if r.font.size else 18) for r in runs)
        text = "".join(r.text for r in runs)
        font = _cjk_font(size)
        if font is not None:
            from PIL import Image, ImageDraw
            d = ImageDraw.Draw(Image.new("RGB", (1, 1)))
            avail = box_w_in * 96 - 8
            lines, cur = 1, ""
            for ch in text:
                if d.textlength(cur + ch, font=font) > avail and cur:
                    lines += 1
                    cur = ch
                else:
                    cur += ch
        else:
            lines = 1
        spacing = p.line_spacing or 1.2
        total_pt += lines * size * spacing * 1.22
        total_pt += (p.space_before.pt if p.space_before else 0) + \
                    (p.space_after.pt if p.space_after else 0)
    return total_pt / 72.0


def text_picture_overlaps(slide):
    """Return the labels of text shapes whose rendered text would run into a picture."""
    from pptx.util import Emu as _E
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    bad = []
    for sh in slide.shapes:
        if sh.shape_type != 17 or not sh.has_text_frame:
            continue
        w_in = sh.width / _E(914400)
        h_in = est_text_height_in(sh, w_in)
        if h_in <= 0:
            continue
        t0 = sh.top / _E(914400)
        r0 = sh.left / _E(914400)
        rect = (r0, t0, r0 + w_in, t0 + h_in)
        for pic in pics:
            pr = (pic.left / _E(914400), pic.top / _E(914400),
                  (pic.left + pic.width) / _E(914400), (pic.top + pic.height) / _E(914400))
            # overlap area (a small tolerance avoids false positives)
            if min(rect[2], pr[2]) - max(rect[0], pr[0]) > 0.05 and \
               min(rect[3], pr[3]) - max(rect[1], pr[1]) > 0.05:
                bad.append(f"{sh.text_frame.text[:24]!r} (height {h_in:.2f} in) -> picture")
    return bad


def audit(path: Path) -> int:
    """font sizes >= 15 pt, shapes inside the canvas, no picture/shape overlap."""
    from pptx.util import Emu as _E
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    problems, runs, smallest = [], 0, 999.0

    for idx, slide in enumerate(prs.slides, 1):
        pics, fills = [], []
        for sh in slide.shapes:
            if sh.left is None:
                continue
            r = (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
            if r[0] < -_E(20000) or r[1] < -_E(20000) or r[2] > W + _E(20000) or r[3] > H + _E(20000):
                problems.append(f"slide {idx}: a shape leaves the canvas")
            if sh.shape_type == 13:                    # PICTURE
                pics.append((r, sh))
            elif sh.shape_type is not None and sh.shape_type != 17:   # not a plain textbox
                try:
                    if sh.fill.type is not None:
                        fills.append((r, sh))
                except Exception:
                    pass
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        runs += 1
                        size = run.font.size.pt if run.font.size else 18
                        smallest = min(smallest, size)
                        if size < MIN_PT - 1e-6:
                            problems.append(f"slide {idx}: {size} pt run {run.text[:24]!r}")
        for pr, psh in pics:
            for fr, fsh in fills:
                if fr[1] < pr[3] - _E(30000) and fr[3] > pr[1] + _E(30000) and \
                   fr[0] < pr[2] - _E(30000) and fr[2] > pr[0] + _E(30000):
                    problems.append(f"slide {idx}: picture overlaps a filled shape "
                                    f"({fsh.text_frame.text[:18]!r})")
        for msg in text_picture_overlaps(slide):
            problems.append(f"slide {idx}: text overlaps a picture - {msg}")

    print(f"slides: {len(prs.slides._sldIdLst)} | runs: {runs} | smallest font: {smallest} pt")
    for p in problems:
        print("FAIL ", p)
    if problems:
        print(f"\nRESULT: {len(problems)} problem(s)")
        return 1
    print("RESULT: OK - every run >= 15 pt, inside the canvas, no overlap")
    return 0


def main() -> int:
    if not FIG.exists() or not list(FIG.glob("*.png")):
        print("figures missing - run python code/make_figures.py first", file=sys.stderr)
        return 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(str(OUT))
    print(f"wrote {OUT}")
    if PROBLEMS:
        for p in PROBLEMS:
            print("WARN ", p)
    return audit(OUT)


if __name__ == "__main__":
    raise SystemExit(main())
